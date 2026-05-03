import os
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, Border, Side, PatternFill
from typing import List, Optional, Callable, Dict


class ExcelFormatter:
    def __init__(self):
        self.progress_callback = None
        self.error_callback = None
        
        self.thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        self.center_alignment = Alignment(horizontal='center', vertical='center')
        self.header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        self.header_font = Font(bold=True, color="FFFFFF")

    def set_progress_callback(self, callback: Callable[[int, str], None]):
        self.progress_callback = callback

    def set_error_callback(self, callback: Callable[[str], None]):
        self.error_callback = callback

    def _report_progress(self, progress: int, message: str):
        if self.progress_callback:
            self.progress_callback(progress, message)

    def _report_error(self, message: str):
        if self.error_callback:
            self.error_callback(message)

    def format_files(
        self,
        file_paths: List[str],
        output_dir: str,
        options: Dict
    ) -> bool:
        if not file_paths:
            self._report_error("未选择任何文件")
            return False
        
        if not os.path.exists(output_dir):
            os.makedirs(output_dir)
        
        total_files = len(file_paths)
        
        try:
            for idx, file_path in enumerate(file_paths):
                self._report_progress(
                    int((idx / total_files) * 90),
                    f"正在处理文件: {os.path.basename(file_path)}"
                )
                
                if not os.path.exists(file_path):
                    self._report_error(f"文件不存在: {file_path}")
                    continue
                
                output_path = os.path.join(
                    output_dir,
                    f"formatted_{os.path.basename(file_path)}"
                )
                
                success = self._format_single_file(
                    file_path,
                    output_path,
                    options
                )
                
                if not success:
                    self._report_error(f"处理文件 {file_path} 失败")
                    continue
            
            self._report_progress(100, f"格式规整完成! 共处理 {total_files} 个文件")
            return True
        
        except Exception as e:
            self._report_error(f"格式规整过程中发生错误: {str(e)}")
            return False

    def _format_single_file(
        self,
        input_path: str,
        output_path: str,
        options: Dict
    ) -> bool:
        try:
            if options.get('use_pandas', False):
                return self._format_with_pandas(input_path, output_path, options)
            else:
                return self._format_with_openpyxl(input_path, output_path, options)
        except Exception as e:
            self._report_error(f"格式化文件 {input_path} 时出错: {str(e)}")
            return False

    def _format_with_pandas(
        self,
        input_path: str,
        output_path: str,
        options: Dict
    ) -> bool:
        df = pd.read_excel(input_path)
        
        if options.get('trim_spaces', True):
            for col in df.columns:
                df[col] = df[col].apply(
                    lambda x: x.strip() if isinstance(x, str) else x
                )
        
        if options.get('remove_duplicates', False):
            df = df.drop_duplicates()
        
        if options.get('fill_missing', False):
            fill_value = options.get('fill_value', '')
            df = df.fillna(fill_value)
        
        df.to_excel(output_path, index=False, engine='openpyxl')
        
        if options.get('apply_styles', True):
            self._apply_styles_to_excel(output_path, options)
        
        return True

    def _format_with_openpyxl(
        self,
        input_path: str,
        output_path: str,
        options: Dict
    ) -> bool:
        wb = load_workbook(input_path)
        
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            
            if options.get('auto_fit_columns', True):
                self._auto_fit_columns(ws)
            
            if options.get('trim_spaces', True):
                self._trim_spaces_in_cells(ws)
            
            if options.get('apply_styles', True):
                self._apply_basic_styles(ws, options)
        
        wb.save(output_path)
        return True

    def _apply_styles_to_excel(self, file_path: str, options: Dict):
        try:
            wb = load_workbook(file_path)
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                
                if options.get('auto_fit_columns', True):
                    self._auto_fit_columns(ws)
                
                if options.get('apply_styles', True):
                    self._apply_basic_styles(ws, options)
            
            wb.save(file_path)
        except Exception as e:
            self._report_error(f"应用样式时出错: {str(e)}")

    def _auto_fit_columns(self, ws):
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

    def _trim_spaces_in_cells(self, ws):
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    cell.value = cell.value.strip()

    def _apply_basic_styles(self, ws, options: Dict):
        for row_idx, row in enumerate(ws.iter_rows(), 1):
            for cell in row:
                if options.get('center_align', True):
                    cell.alignment = self.center_alignment
                
                if options.get('add_borders', True):
                    cell.border = self.thin_border
                
                if options.get('style_header', True) and row_idx == 1:
                    cell.font = self.header_font
                    cell.fill = self.header_fill

    def quick_format(
        self,
        file_path: str,
        output_path: str
    ) -> bool:
        default_options = {
            'trim_spaces': True,
            'center_align': True,
            'add_borders': True,
            'auto_fit_columns': True,
            'style_header': True,
            'apply_styles': True,
            'use_pandas': False
        }
        return self._format_single_file(file_path, output_path, default_options)
